# SESHAT — Egyptian goddess of writing, record-keeping, and the House of
# Life (the temple library). She records deeds onto the leaves of the sacred
# tree so they are never lost. SESHAT is ODIN's librarian: hand her any link
# — webpage, GitHub repo, YouTube video, image, PDF — and she extracts its
# substance, has SARASWATI distill it into a knowledge note (cloud calls stay
# inside SARASWATI per the cloud-AI policy), writes the note to the Obsidian
# vault, and from then on ODIN answers questions about it fully offline.
#
# Degrade path (offline / all providers down): the RAW capture is still
# saved to the vault and the topic is pushed onto SARASWATI's research queue
# so SELENE's idle drain distills it later. Learning is never silently lost.

import os
import re
import json
from datetime import datetime

from core.marduk import OdinModule

try:
    import requests
    _HAS_REQUESTS = True
except ImportError:
    _HAS_REQUESTS = False

# Optional — only needed for YouTube transcripts. pip install youtube-transcript-api
try:
    from youtube_transcript_api import YouTubeTranscriptApi
    _HAS_YT = True
except ImportError:
    _HAS_YT = False

_IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp")


def _slugify(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", (s or "").lower()).strip("_")[:60] or "untitled"
_GITHUB_RE = re.compile(r"^https?://(?:www\.)?github\.com/([\w.\-]+)/([\w.\-]+?)(?:\.git)?(?:[/#?].*)?$", re.I)
_YT_RE = re.compile(r"^https?://(?:www\.)?(?:youtube\.com/watch\?.*v=([\w\-]{6,})|youtu\.be/([\w\-]{6,}))", re.I)


class Seshat(OdinModule):
    MODULE_NAME = "SESHAT"
    LAYER = "INTELLIGENCE"

    def __init__(self, config: dict):
        super().__init__(config)
        cfg = config.get("seshat", {})
        self.download_dir = cfg.get("download_dir", "data/seshat")
        self.max_extract_chars = int(cfg.get("max_extract_chars", 12000))
        self.max_raw_save_chars = int(cfg.get("max_raw_save_chars", 20000))
        # docling = high-fidelity local doc extraction (tables/layout/OCR) but
        # heavy + slow cold start, so off by default. Light libs handle the
        # common case instantly and offline. See _extract_local_file.
        self.use_docling = bool(cfg.get("use_docling", False))
        self._docling_converter = None
        self.vault_root = os.path.expanduser(config.get("nabu", {}).get("vault_root", "~/Brain"))
        # Working memory of the last thing learned (link or file), so follow-ups
        # like "summarize that" / "what did it say about X" answer instantly.
        self._last_learned_path = "data/knowledge/last_learned.json"

    @property
    def skills(self) -> list[dict]:
        return [
            {
                "name": "learn_link",
                "description": (
                    "Teach ODIN from a URL. Works on webpages/articles, GitHub repos, "
                    "YouTube videos, images/photos, and PDFs. Extracts the content, "
                    "distills it into a knowledge note, saves it to the vault so future "
                    "questions on the topic answer offline. Use when the user shares a "
                    "link and says 'learn this' / 'study this' / 'read this'."
                ),
                "parameters": {
                    "url":   {"type": "string", "description": "The link to learn from"},
                    "focus": {"type": "string", "description": "Optional: what specifically to learn from it"},
                },
                "required": ["url"],
            },
            {
                "name": "learn_file",
                "description": (
                    "Read a document FILE on this computer and learn it — PDF, Word, "
                    "PowerPoint, text, code, CSV. Extracts the content, distills it "
                    "into a vault note, and answers about it offline thereafter. Use "
                    "when the user points at a file on disk (not a URL)."
                ),
                "parameters": {
                    "path":  {"type": "string", "description": "File path OR just a filename — ODIN searches Desktop/Downloads/Documents/vault if it's not a full path."},
                    "focus": {"type": "string", "description": "Optional: what to focus on."},
                },
                "required": ["path"],
            },
            {
                "name": "note_context",
                "description": (
                    "Internal: record content into ODIN's working memory as 'the "
                    "last thing ingested', so follow-ups like 'what did it say "
                    "about X' work. Used by other modules (e.g. HORUS pushes what "
                    "it just saw on screen)."
                ),
                "parameters": {
                    "title":   {"type": "string", "description": "Short label for what this is."},
                    "content": {"type": "string", "description": "The text content to remember."},
                    "kind":    {"type": "string", "description": "Source kind (e.g. 'screen', 'webpage')."},
                },
                "required": ["title", "content"],
                "internal_only": True,
            },
            {
                "name": "recall_learned",
                "description": (
                    "Answer a follow-up about the LAST thing ODIN learned (the "
                    "link or file it just ingested). Use for 'summarize that', "
                    "'what did it say about X', 'what was in it', 'recap that'. "
                    "Answers from the saved note, so it's instant and works offline."
                ),
                "parameters": {
                    "question": {"type": "string", "description": "Optional specific question; empty = summarize."},
                },
                "required": [],
            },
        ]

    def execute(self, skill_name: str, args: dict) -> str:
        if skill_name == "learn_link":
            return self._learn_link(args.get("url", ""), args.get("focus", ""))
        if skill_name == "learn_file":
            return self._learn_file(args.get("path", ""), args.get("focus", ""))
        if skill_name == "recall_learned":
            return self._recall_learned(args.get("question", ""))
        if skill_name == "note_context":
            title = args.get("title", "") or "untitled"
            content = args.get("content", "")
            if not content.strip():
                return "Nothing to note."
            self._record_last_learned(title, title, args.get("kind", "context"),
                                      content, _slugify(title))
            return f"Noted: {title}"
        return f"[SESHAT] Unknown skill: {skill_name}"

    # === Main flow =======================================================

    def _learn_link(self, url: str, focus: str = "") -> str:
        url = (url or "").strip().rstrip(".,;!?")
        if not url.lower().startswith(("http://", "https://")):
            return "I need a full link starting with http or https."

        kind, title, content = self._extract(url)
        if not content or not content.strip():
            return (f"I couldn't pull anything readable from that {kind or 'link'}. "
                    f"If it needs a login or heavy JavaScript, try ARGUS browsing instead.")
        return self._distill_and_save(kind, title or url, content, url, focus)

    def _learn_file(self, path: str, focus: str = "") -> str:
        resolved = self._resolve_file(path)
        if not resolved:
            return (f"I couldn't find '{os.path.basename((path or '').strip()) or path}' "
                    f"on the Desktop, Downloads, Documents, or the vault. "
                    f"Give me the full path if it's somewhere else.")
        path = resolved
        kind, title, content = self._extract_local_file(path)
        if not content or not content.strip():
            return (f"I couldn't pull readable text out of {os.path.basename(path)}. "
                    f"If it's a scanned PDF or image-heavy, enable seshat.use_docling for OCR.")
        return self._distill_and_save(kind, title or os.path.basename(path),
                                      content, os.path.basename(path), focus)

    def _distill_and_save(self, kind: str, title: str, content: str,
                          source: str, focus: str = "") -> str:
        content = content[: self.max_extract_chars]
        title = title or source

        # Distill via SARASWATI (the only module allowed to spend cloud calls).
        note_md, provider = "", ""
        sara = self.marduk.get_module("SARASWATI") if self.marduk else None
        if sara is not None:
            try:
                note_md, provider = sara.distill(title, content, focus)
            except Exception as e:
                print(f"[SESHAT] distill failed: {e}")

        slug = re.sub(r"[^a-z0-9]+", "_", title.lower()).strip("_")[:60] or "untitled"
        ts = datetime.now().strftime("%Y-%m-%d %H:%M")

        if note_md:
            body = (f"# {title}\n\n> learned {ts} from {source} ({kind}, distilled via {provider})\n\n"
                    f"{note_md}\n")
            status = f"Learned and archived. Ask me about it anytime — even offline."
        else:
            # Offline / providers exhausted: keep the raw capture (still
            # searchable by NABU/HERMES) and queue the topic so SELENE's
            # idle research drain distills it when the cloud comes back.
            body = (f"# {title}\n\n> raw capture {ts} from {source} ({kind}) — distill pending\n\n"
                    f"{content[: self.max_raw_save_chars]}\n")
            status = ("Cloud is unreachable, so I archived the raw content and queued it "
                      "for distillation when I'm back online. It's already searchable.")
            if sara is not None:
                try:
                    sara.execute("queue_research", {"topic": f"{title} — {source}"})
                except Exception:
                    pass

        saved = self._save_note(f"learned/{slug}.md", body)
        if not saved:
            return "I extracted the content but couldn't write to the vault — check NABU."

        # Remember this as the 'last learned' so follow-ups ("summarize that",
        # "what did it say about X") answer instantly from the distilled note.
        self._record_last_learned(title, source, kind, note_md or content[:4000], slug)

        # First two sentences of the note as the spoken gist.
        gist = " ".join(re.split(r"(?<=[.!?])\s+", re.sub(r"[#>*`]", "", note_md or content).strip())[:2])
        return f"{status} Saved as learned/{slug}. {gist[:300]}"

    # === Working memory: last learned + follow-up recall =================
    def _record_last_learned(self, title, source, kind, note_text, slug):
        try:
            os.makedirs(os.path.dirname(self._last_learned_path), exist_ok=True)
            with open(self._last_learned_path, "w", encoding="utf-8") as f:
                json.dump({"title": title, "source": source, "kind": kind,
                           "slug": slug, "note": note_text[:6000],
                           "ts": datetime.now().isoformat(timespec="seconds")}, f)
        except OSError as e:
            print(f"[SESHAT] last-learned write failed: {e}")

    def _recall_learned(self, question: str = "") -> str:
        if not os.path.exists(self._last_learned_path):
            return "I haven't learned anything yet this session — give me a link or a file first."
        try:
            with open(self._last_learned_path, "r", encoding="utf-8") as f:
                last = json.load(f)
        except (OSError, json.JSONDecodeError):
            return "I can't recall the last thing I learned — the note seems lost."
        title, note = last.get("title", "that"), last.get("note", "")
        if not note.strip():
            return f"I learned '{title}' but kept no readable note of it."
        # Bridge: if what's in mind is a FOUND PATH (from locate_file /
        # find_files_smart), "summarize it" should read the actual file —
        # learn it now (which re-records working memory as the document, so
        # further follow-ups interrogate its content).
        if last.get("kind") in ("file-found", "files-found"):
            for line in note.splitlines():
                p = line.strip()
                if p and os.path.isfile(p):
                    return self._learn_file(p, focus=question)
            return (f"I found '{title}' but it isn't a readable file "
                    f"(a folder, maybe). Tell me exactly what to do with it.")
        q = (question or "").strip()
        if not q:
            # Summarize: the distilled note already leads with a summary.
            gist = re.sub(r"^#.*$", "", note, flags=re.M).strip()
            return f"From '{title}': {gist[:600]}"
        # Specific question — answer from the note (cloud if available, cheap
        # since the note is already distilled; else keyword-grep the note).
        sara = self.marduk.get_module("SARASWATI") if self.marduk else None
        if sara is not None and getattr(sara, "providers", None):
            try:
                out = sara.execute("ask_ai", {"question":
                    f"Using only this note about '{title}', answer concisely.\n\n"
                    f"NOTE:\n{note}\n\nQUESTION: {q}"})
                if out and not str(out).lower().startswith(("[", "marduk")):
                    return str(out)
            except Exception:
                pass
        # Offline grep: return note lines mentioning the question's keywords.
        kws = [w for w in re.findall(r"\w+", q.lower()) if len(w) > 3]
        hits = [ln.strip() for ln in note.splitlines()
                if ln.strip() and any(w in ln.lower() for w in kws)]
        if hits:
            return f"From '{title}': " + " ".join(hits[:4])[:600]
        return (f"'{title}' doesn't seem to mention that. Here's the gist instead: "
                f"{re.sub(r'^#.*$', '', note, flags=re.M).strip()[:400]}")

    def _resolve_file(self, path: str) -> str:
        """Turn a path-or-bare-filename into a real file. Exact path wins;
        otherwise search the user's common folders + the vault for a matching
        filename so 'summarize report.pdf' finds report.pdf without a full path."""
        path = (path or "").strip().strip('"').strip("'")
        if not path:
            return ""
        if os.path.exists(path):
            return path
        exp = os.path.expanduser(os.path.expandvars(path))
        if os.path.exists(exp):
            return exp
        name = os.path.basename(path).lower()
        home = os.path.expanduser("~")
        roots = [os.getcwd(),
                 os.path.join(home, "Desktop"), os.path.join(home, "Downloads"),
                 os.path.join(home, "Documents"),
                 os.path.join(home, "OneDrive", "Desktop"),
                 os.path.join(home, "OneDrive", "Documents"),
                 self.vault_root]
        for root in roots:
            if not os.path.isdir(root):
                continue
            cand = os.path.join(root, os.path.basename(path))
            if os.path.exists(cand):
                return cand
            try:
                for entry in os.scandir(root):
                    if entry.is_file() and entry.name.lower() == name:
                        return entry.path
            except OSError:
                pass
            # one level into subfolders (covers Desktop/ProjectX/report.pdf)
            try:
                for entry in os.scandir(root):
                    if entry.is_dir():
                        sub = os.path.join(entry.path, os.path.basename(path))
                        if os.path.exists(sub):
                            return sub
            except OSError:
                pass
        return ""

    # === Extraction by link type ========================================

    def _extract(self, url: str) -> tuple[str, str, str]:
        """Returns (kind, title, content). Empty content = extraction failed."""
        m = _GITHUB_RE.match(url)
        if m:
            return self._extract_github(m.group(1), m.group(2))
        m = _YT_RE.match(url)
        if m:
            return self._extract_youtube(url, m.group(1) or m.group(2))
        path = url.split("?", 1)[0].lower()
        if path.endswith(_IMAGE_EXTS):
            return self._extract_image(url)
        if path.endswith(".pdf"):
            return self._extract_pdf(url)
        return self._extract_page(url)

    def _extract_github(self, owner: str, repo: str) -> tuple[str, str, str]:
        """Public GitHub API — no key, no clone. README + metadata + file map
        is plenty for a knowledge note; full code reading is a CHITRA/clone job."""
        if not _HAS_REQUESTS:
            return ("repo", f"{owner}/{repo}", "")
        parts = []
        try:
            meta = requests.get(f"https://api.github.com/repos/{owner}/{repo}", timeout=10).json()
            parts.append(
                f"Repo: {meta.get('full_name')}\nDescription: {meta.get('description')}\n"
                f"Language: {meta.get('language')}  Stars: {meta.get('stargazers_count')}\n"
                f"Topics: {', '.join(meta.get('topics') or [])}"
            )
            title = meta.get("full_name") or f"{owner}/{repo}"
        except Exception as e:
            print(f"[SESHAT] github meta failed: {e}")
            title = f"{owner}/{repo}"
        try:
            files = requests.get(f"https://api.github.com/repos/{owner}/{repo}/contents", timeout=10).json()
            if isinstance(files, list):
                parts.append("Top-level files: " + ", ".join(f.get("name", "?") for f in files[:40]))
        except Exception:
            pass
        try:
            r = requests.get(f"https://api.github.com/repos/{owner}/{repo}/readme",
                             headers={"Accept": "application/vnd.github.raw"}, timeout=10)
            if r.ok:
                parts.append("README:\n" + r.text)
        except Exception as e:
            print(f"[SESHAT] github readme failed: {e}")
        return ("github repo", title, "\n\n".join(parts))

    def _extract_youtube(self, url: str, video_id: str) -> tuple[str, str, str]:
        title = ""
        if _HAS_REQUESTS:
            try:
                o = requests.get("https://www.youtube.com/oembed",
                                 params={"url": url, "format": "json"}, timeout=10).json()
                title = o.get("title", "")
            except Exception:
                pass
        transcript = ""
        if _HAS_YT:
            try:
                # 0.6.x API first, 1.x API as fallback.
                try:
                    entries = YouTubeTranscriptApi.get_transcript(video_id, languages=["en", "hi"])
                    transcript = " ".join(e["text"] for e in entries)
                except AttributeError:
                    fetched = YouTubeTranscriptApi().fetch(video_id, languages=["en", "hi"])
                    transcript = " ".join(s.text for s in fetched)
            except Exception as e:
                print(f"[SESHAT] transcript failed: {e}")
        else:
            print("[SESHAT] youtube-transcript-api not installed — pip install youtube-transcript-api")
        if not transcript:
            # No transcript → at least learn from the watch-page description.
            _, t2, page = self._extract_page(url)
            return ("youtube video", title or t2, page)
        return ("youtube video", title or url, f"Video title: {title}\nTranscript:\n{transcript}")

    def _extract_image(self, url: str) -> tuple[str, str, str]:
        if not _HAS_REQUESTS:
            return ("image", url, "")
        os.makedirs(self.download_dir, exist_ok=True)
        name = re.sub(r"[^\w.\-]", "_", url.rsplit("/", 1)[-1])[:80] or "image.png"
        local = os.path.join(self.download_dir, name)
        try:
            r = requests.get(url, timeout=20)
            r.raise_for_status()
            with open(local, "wb") as f:
                f.write(r.content)
        except Exception as e:
            print(f"[SESHAT] image download failed: {e}")
            return ("image", url, "")
        parts = [f"Image file: {name} (saved to {local})"]
        horus = self.marduk.get_module("HORUS") if self.marduk else None
        if horus:
            try:
                parts.append("Objects detected: " + horus.execute("detect_objects", {"path": local}))
            except Exception:
                pass
        try:
            import pytesseract
            from PIL import Image
            text = pytesseract.image_to_string(Image.open(local)).strip()
            if text:
                parts.append("Text in image (OCR):\n" + text)
        except Exception:
            pass
        return ("image", name, "\n\n".join(parts))

    def _extract_pdf(self, url: str) -> tuple[str, str, str]:
        if not _HAS_REQUESTS:
            return ("pdf", url, "")
        os.makedirs(self.download_dir, exist_ok=True)
        name = re.sub(r"[^\w.\-]", "_", url.rsplit("/", 1)[-1])[:80] or "doc.pdf"
        local = os.path.join(self.download_dir, name)
        try:
            r = requests.get(url, timeout=30)
            r.raise_for_status()
            with open(local, "wb") as f:
                f.write(r.content)
        except Exception as e:
            print(f"[SESHAT] pdf download failed: {e}")
            return ("pdf", url, "")
        stirling = self.marduk.get_module("STIRLING") if self.marduk else None
        if stirling:
            try:
                text = stirling.execute("pdf_to_text", {"path": local, "max_chars": self.max_extract_chars})
                return ("pdf", name, text)
            except Exception as e:
                print(f"[SESHAT] pdf_to_text failed: {e}")
        return ("pdf", name, "")

    # === Local file extraction (docling-preferred, light fallback) =======
    # docling gives structure/table/OCR fidelity but is heavy (~90s cold
    # first call, big RAM) so it's OFF by default and lazy-imported only
    # when seshat.use_docling is true. The light libs (pdfplumber/python-
    # docx/python-pptx) are instant, offline, and need no model downloads,
    # so they're the default and the fallback.
    _TEXT_EXTS = (".txt", ".md", ".markdown", ".csv", ".log", ".json", ".yaml", ".yml",
                  ".py", ".js", ".ts", ".java", ".c", ".cpp", ".cs", ".go", ".rs",
                  ".rb", ".php", ".sh", ".ps1", ".sql", ".html", ".css", ".xml")

    def _extract_local_file(self, path: str) -> tuple[str, str, str]:
        ext = os.path.splitext(path)[1].lower()
        name = os.path.basename(path)

        # Images reuse the existing OCR/vision path.
        if ext in _IMAGE_EXTS:
            return self._extract_image_local(path)

        # High-fidelity path: docling handles pdf/docx/pptx/html/md with
        # layout + tables + (optional) OCR.
        if self.use_docling and ext in (".pdf", ".docx", ".pptx", ".html", ".htm", ".md"):
            md = self._docling_markdown(path)
            if md:
                return (f"document ({ext.lstrip('.')})", name, md)
            # fall through to the light path if docling failed

        if ext == ".pdf":
            # Reuse STIRLING's pdf_to_text, else pdfplumber.
            stirling = self.marduk.get_module("STIRLING") if self.marduk else None
            if stirling:
                try:
                    text = stirling.execute("pdf_to_text", {"path": path, "max_chars": self.max_extract_chars})
                    if text and text.strip():
                        return ("pdf", name, text)
                except Exception as e:
                    print(f"[SESHAT] STIRLING pdf_to_text failed: {e}")
            return ("pdf", name, self._pdfplumber_text(path))

        if ext == ".docx":
            return ("word document", name, self._docx_text(path))
        if ext == ".pptx":
            return ("powerpoint", name, self._pptx_text(path))
        if ext in self._TEXT_EXTS:
            try:
                with open(path, "r", encoding="utf-8", errors="replace") as f:
                    return ("text file", name, f.read())
            except Exception as e:
                return ("text file", name, "")
        # Unknown binary — last try: read as text.
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                return ("file", name, f.read())
        except Exception:
            return ("file", name, "")

    def _docling_markdown(self, path: str) -> str:
        try:
            if self._docling_converter is None:
                from docling.document_converter import DocumentConverter
                self._docling_converter = DocumentConverter()
            return self._docling_converter.convert(path).document.export_to_markdown()
        except Exception as e:
            print(f"[SESHAT] docling failed ({e}); using light extractor.")
            return ""

    def _pdfplumber_text(self, path: str) -> str:
        try:
            import pdfplumber
            out = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    out.append(page.extract_text() or "")
                    if sum(len(p) for p in out) > self.max_extract_chars:
                        break
            return "\n\n".join(out).strip()
        except Exception as e:
            print(f"[SESHAT] pdfplumber failed: {e}")
            return ""

    def _docx_text(self, path: str) -> str:
        try:
            import docx
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts).strip()
        except Exception as e:
            print(f"[SESHAT] python-docx failed: {e}")
            return ""

    def _pptx_text(self, path: str) -> str:
        try:
            from pptx import Presentation
            parts = []
            for i, slide in enumerate(Presentation(path).slides, 1):
                parts.append(f"--- Slide {i} ---")
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        parts.append(shape.text_frame.text.strip())
            return "\n".join(parts).strip()
        except Exception as e:
            print(f"[SESHAT] python-pptx failed: {e}")
            return ""

    def _extract_image_local(self, path: str) -> tuple[str, str, str]:
        name = os.path.basename(path)
        parts = [f"Image file: {name}"]
        horus = self.marduk.get_module("HORUS") if self.marduk else None
        if horus:
            try:
                parts.append("Objects detected: " + horus.execute("detect_objects", {"path": path}))
            except Exception:
                pass
        try:
            import pytesseract
            from PIL import Image
            text = pytesseract.image_to_string(Image.open(path)).strip()
            if text:
                parts.append("Text in image (OCR):\n" + text)
        except Exception:
            pass
        return ("image", name, "\n\n".join(parts))

    def _extract_page(self, url: str) -> tuple[str, str, str]:
        akasha = self.marduk.get_module("AKASHA") if self.marduk else None
        if akasha:
            try:
                md = akasha.execute("fetch_page", {"url": url})
                if md and not md.lower().startswith(("[akasha]", "error", "failed", "could not")):
                    # First markdown heading doubles as the title.
                    m = re.search(r"^#\s+(.+)$", md, re.M)
                    return ("webpage", (m.group(1).strip() if m else url), md)
            except Exception as e:
                print(f"[SESHAT] fetch_page failed: {e}")
        return ("webpage", url, "")

    # === Vault ===========================================================

    def _save_note(self, rel_path: str, body: str) -> bool:
        nabu = self.marduk.get_module("NABU") if self.marduk else None
        if not nabu:
            return False
        try:
            res = nabu.execute("write_note", {"path": rel_path, "content": body, "append": False})
            return bool(res) and "fail" not in str(res).lower()
        except Exception as e:
            print(f"[SESHAT] vault write failed: {e}")
            return False
